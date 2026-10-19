# -*- coding: utf-8 -*-
"""تشخيص ملف الأسعار"""
import sys
from pathlib import Path

# اقرأ المسار من المستخدم
if len(sys.argv) > 1:
    file_path = sys.argv[1]
else:
    file_path = input("اسحب ملف الأسعار هنا ثم اضغط Enter: ").strip('"').strip("'")

p = Path(file_path)
if not p.exists():
    print("[!] الملف غير موجود: " + str(p))
    sys.exit(1)

print("=" * 60)
print("Diagnosing: " + str(p))
print("=" * 60)
print("Extension: " + p.suffix)
print("Size: " + str(p.stat().st_size) + " bytes")
print()

raw = p.read_bytes()
print("First 8 bytes (hex):", ' '.join(f'{b:02X}' for b in raw[:8]))
if raw[:4] == b'PK\x03\x04':
    print("  → Looks like .xlsx (ZIP)")
elif raw[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
    print("  → Looks like .xls (OLE2 / old Excel)")
else:
    print("  → Unknown format")
print()

# قراءة بـ openpyxl
try:
    import openpyxl
    from io import BytesIO
    print("Trying openpyxl...")
    wb = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    print("  Sheets: " + str(wb.sheetnames))
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        print()
        print("--- Sheet: " + sheet + " (" + str(len(rows)) + " rows) ---")
        for i, row in enumerate(rows[:5]):
            print("  Row " + str(i) + ": " + repr(row))
except Exception as e:
    print("[!] openpyxl failed: " + str(e))
    print()
    print("Trying with pandas + xlrd for .xls...")
    try:
        import pandas as pd
        xls = pd.ExcelFile(BytesIO(raw))
        for sheet in xls.sheet_names:
            df = pd.read_excel(BytesIO(raw), sheet_name=sheet)
            print("--- Sheet: " + sheet + " ---")
            print(df.head(5))
    except Exception as e2:
        print("[!] pandas also failed: " + str(e2))

# محاكاة الـ parser
print()
print("=" * 60)
print("Simulating parser...")
print("=" * 60)
try:
    wb = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            print("Sheet '" + sheet + "': <2 rows, skipped")
            continue
        header = [str(c).lower().strip() if c else '' for c in rows[0]]
        print("Sheet '" + sheet + "' headers: " + str(header))
        col = {}
        for i, h in enumerate(header):
            if 'system' in h or 'نظام' in h:
                col['system'] = i
                print("  → system at col " + str(i))
            elif 'size' in h or 'قطر' in h or 'diameter' in h:
                col['size'] = i
                print("  → size at col " + str(i))
            elif 'supply' in h or 'توريد' in h:
                col['supply'] = i
                print("  → supply at col " + str(i))
            elif 'install' in h or 'تركيب' in h:
                col['install'] = i
                print("  → install at col " + str(i))
            elif 'test' in h or 'اختبار' in h:
                col['test'] = i
                print("  → test at col " + str(i))
        print("Cols detected: " + str(col))
        if 'system' not in col or 'size' not in col:
            print("  [!] Missing system or size → skipped")
        else:
            print("  [+] Valid sheet")
except Exception as e:
    print("[!] " + str(e))